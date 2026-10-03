import io

import pytest
from unibot_RAG.config import Settings
from unibot_RAG.ingestion.loaders import load_documents
from unibot_RAG.ingestion.chunking import chunk_documents, token_count


def test_invalid_overlap():
    with pytest.raises(ValueError):
        Settings(chunk_tokens=20, chunk_overlap=20)


@pytest.mark.parametrize(
    "filename,content",
    [
        ("a.txt", b"hello world"),
        ("a.md", b"# Heading\n\nA paragraph."),
        ("a.html", b"<h1>Title</h1><script>evil()</script><p>body</p>"),
        ("a.csv", b"item,cost\nBolt,12\nNut,6"),
        ("a.json", b'[{"name":"Bolt","cost":12}]'),
        ("a.jsonl", b'{"name":"Bolt"}\n{"name":"Nut"}'),
    ],
)
def test_load_formats(filename, content):
    docs = load_documents(filename, content, "source-a", Settings())
    assert docs and all(d.source_id == "source-a" and d.text.strip() for d in docs)
    assert "evil()" not in "\n".join(d.text for d in docs)


def test_faq_keeps_answer_and_context():
    docs = load_documents(
        "faq.csv",
        b"Question,Answer,Context,uuid\nHow?,Carefully,Manual,a1",
        "faq",
        Settings(),
    )
    assert all(word in docs[0].text for word in ["How?", "Carefully", "Manual"])
    assert docs[0].metadata["record"] == 1


@pytest.mark.parametrize(
    "filename,content",
    [
        ("x.json", b"null"),
        ("x.csv", b"a,b\n1,2,3"),
        ("x.txt", b"  "),
        ("x.exe", b"hello"),
    ],
)
def test_bad_input_rejected(filename, content):
    with pytest.raises(ValueError):
        load_documents(filename, content, "s", Settings())


@pytest.mark.parametrize("strategy", ["token", "recursive", "structure"])
def test_chunks_bounded_provenance_and_stable_ids(strategy):
    cfg = Settings(chunk_strategy=strategy, chunk_tokens=40, chunk_overlap=5)
    text = "# 採購\n\n" + "零件價格為十二元。注意庫存數量。" * 50
    docs = load_documents("x.md", text.encode(), "s", cfg)
    chunks = chunk_documents(docs, cfg)
    assert len(chunks) > 2
    assert all(token_count(c.text) <= 40 for c in chunks)
    assert all(c.source_id == "s" and "\ufffd" not in c.text for c in chunks)
    assert [c.id for c in chunks] == [c.id for c in chunk_documents(docs, cfg)]


def test_docx_loader():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Shipping instructions")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Part"
    table.rows[0].cells[1].text = "Cost"
    output = io.BytesIO()
    doc.save(output)
    docs = load_documents("x.docx", output.getvalue(), "s", Settings())
    assert "Shipping instructions" in docs[0].text
    assert "Part" in docs[0].text


def test_semantic_requires_embeddings_and_preserves_text():
    cfg = Settings(chunk_strategy="semantic", chunk_tokens=40, chunk_overlap=0)
    docs = load_documents("x.txt", b"First topic.\n\nSecond topic.", "s", cfg)
    with pytest.raises(ValueError):
        chunk_documents(docs, cfg)
    chunks = chunk_documents(docs, cfg, embed=lambda texts: [[1.0, 0.0], [0.0, 1.0]])
    assert len(chunks) == 2
    assert "First topic." in chunks[0].text


def test_recursive_packs_short_paragraphs_before_splitting():
    cfg = Settings(chunk_strategy="recursive", chunk_tokens=40, chunk_overlap=4)
    docs = load_documents(
        "a.txt", b"First short paragraph.\n\nSecond short paragraph.", "s", cfg
    )
    chunks = chunk_documents(docs, cfg)
    assert len(chunks) == 1
    assert "First" in chunks[0].text and "Second" in chunks[0].text


def test_chunk_limit_stops_before_exhausting_a_huge_document():
    cfg = Settings(chunk_tokens=16, chunk_overlap=15, max_chunks=2)
    docs = load_documents("a.txt", b"word " * 10000, "s", cfg)
    with pytest.raises(ValueError, match="Chunk limit"):
        chunk_documents(docs, cfg)
