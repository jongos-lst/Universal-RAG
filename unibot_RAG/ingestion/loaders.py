"""Pure byte-to-document adapters; no arbitrary filesystem or network reads."""

import csv
import io
import json
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup
from unibot_RAG.config import Settings
from unibot_RAG.domain import Document

SUPPORTED = {".txt", ".md", ".html", ".htm", ".csv", ".json", ".jsonl", ".pdf", ".docx"}


def html_text(text: str) -> str:
    soup = BeautifulSoup(text, "html.parser")
    for node in soup(["script", "style", "nav", "footer", "header"]):
        node.decompose()
    for heading in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6"]):
        heading.replace_with(
            "\n"
            + "#" * int(heading.name[1])
            + " "
            + heading.get_text(" ", strip=True)
            + "\n"
        )
    return soup.get_text("\n", strip=True)


def load_documents(
    filename: str, content: bytes, source_id: str, settings: Settings
) -> list[Document]:
    suffix = Path(filename).suffix.lower()
    if (
        suffix not in SUPPORTED
        or not content
        or len(content) > settings.max_input_bytes
    ):
        raise ValueError(
            "Unsupported format, empty input, or input byte limit exceeded"
        )
    if not source_id.strip() or len(source_id) > 1024:
        raise ValueError("Invalid source_id")
    records = []
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > settings.max_records:
            raise ValueError("Encrypted PDF or page limit exceeded")
        for page, item in enumerate(reader.pages, 1):
            records.append((item.extract_text() or "", {"page": page}))
    elif suffix == ".docx":
        from docx import Document as WordDocument

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if (
                sum(item.file_size for item in archive.infolist())
                > settings.max_input_bytes * 10
            ):
                raise ValueError("DOCX expansion limit exceeded")
        doc = WordDocument(io.BytesIO(content))
        text = "\n\n".join(p.text for p in doc.paragraphs)
        tables = [
            "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            for table in doc.tables
        ]
        records = [(text + "\n\n" + "\n\n".join(tables), {})]
    else:
        text = content.decode("utf-8-sig")
        if suffix in {".html", ".htm"}:
            records = [(html_text(text), {})]
        elif suffix == ".csv":
            reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames or len(set(reader.fieldnames)) != len(
                reader.fieldnames
            ):
                raise ValueError("CSV needs unique column names")
            for number, row in enumerate(reader, 1):
                if (
                    number > settings.max_records
                    or None in row
                    or any(v is None for v in row.values())
                ):
                    raise ValueError("CSV row shape or record limit invalid")
                records.append(
                    ("\n".join(f"{k}: {v}" for k, v in row.items()), {"record": number})
                )
        elif suffix in {".json", ".jsonl"}:
            values = (
                [json.loads(line) for line in text.splitlines() if line.strip()]
                if suffix == ".jsonl"
                else json.loads(text)
            )
            if isinstance(values, dict):
                values = [values]
            if not isinstance(values, list) or len(values) > settings.max_records:
                raise ValueError("JSON must be an object or bounded array of objects")
            for number, row in enumerate(values, 1):
                if not isinstance(row, dict) or not row:
                    raise ValueError("JSON records must be nonempty objects")
                records.append(
                    (
                        json.dumps(row, ensure_ascii=False, sort_keys=True),
                        {"record": number},
                    )
                )
        else:
            records = [(text, {})]
    if len(records) > settings.max_records:
        raise ValueError("Record limit exceeded")
    docs = [
        Document(
            text=text.strip(),
            source_id=source_id,
            metadata={"filename": Path(filename).name, "format": suffix[1:], **meta},
        )
        for text, meta in records
        if text.strip()
    ]
    if (
        not docs
        or sum(len(d.text.encode()) for d in docs) > settings.max_input_bytes * 10
    ):
        raise ValueError(
            "No extractable text (OCR is not supported), or expanded text limit exceeded"
        )
    return docs
